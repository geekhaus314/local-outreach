"""Branded outreach ledger: leads.csv -> leads_ledger.xlsx.

Real accounting/marketing fields only. Nothing is invented here -- every cell
ancored to a measured column in leads.csv or derived from it.
"""

from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "leads.csv"
OUT_PATH = ROOT / "leads_ledger.xlsx"

BRAND_NAME = "Jake Viefhaus"
BRAND_SITE = "webrehabstl.pages.dev"
BRAND_LOC = "Florissant, MO"

INK = "1A1A2E"
PAPER = "FFFFFF"
ACCENT = "E8622A"
MUTED = "6B7280"
BANDFILL = "F7F7F5"

thin = Side(style="thin", color="D9D9D6")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_cell(c, bold=False, color=INK, size=10, fill=None, wrap=False, halign=None):
    c.font = Font(name="Helvetica", size=size, bold=bold, color=color)
    c.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal=halign)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)


def brand_block(ws, title: str, subtitle: str):
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{BRAND_NAME}  ·  {BRAND_SITE}  ·  {BRAND_LOC}"
    style_cell(ws["A1"], bold=True, size=11, color=ACCENT)
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[1].height = 22
    ws.merge_cells("A2:H2")
    ws["A2"] = title
    style_cell(ws["A2"], bold=True, size=16)
    ws["A2"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 30
    ws.merge_cells("A3:H3")
    ws["A3"] = subtitle
    style_cell(ws["A3"], size=10, color=MUTED)
    ws["A3"].alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[3].height = 18
    for col in range(1, 9):
        ws.cell(row=4, column=col).fill = PatternFill("solid", fgColor=ACCENT)
    ws.row_dimensions[4].height = 3


def table(ws, headers, rows, start_row, widths, band=True):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        style_cell(c, bold=True, size=9, color=PAPER, fill=INK, wrap=True)
        c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left")
        c.border = BRD
    ws.row_dimensions[start_row].height = 28
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        fill = BANDFILL if band and i % 2 else None
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            style_cell(c, size=9, fill=fill, wrap=True)
            c.border = BRD
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 30
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"A{start_row + 1}"


def num(s) -> float:
    try:
        return float(s)
    except (TypeError, ValueError):
        return 0.0


def main():
    rows = list(csv.DictReader(CSV_PATH.open(newline="", encoding="utf-8")))
    wb = Workbook()

    # --- Sheet 1: dashboard -------------------------------------------------
    dash = wb.active
    if dash is None:
        raise RuntimeError('no active sheet')
    dash.title = "Dashboard"
    dash.sheet_view.showGridLines = False
    brand_block(
        dash,
        "Outreach Ledger",
        f"Generated {date.today().isoformat()} from leads.csv. "
        f"{len(rows)} businesses on file.",
    )
    by_cls: dict[str, int] = {}
    by_src: dict[str, int] = {}
    opps = []
    for r in rows:
        by_cls[r["Classification"]] = by_cls.get(r["Classification"], 0) + 1
        by_src[r["Sources"]] = by_src.get(r["Sources"], 0) + 1
        opps.append(num(r["Opportunity Score"]))
    stats = [
        ("Total leads", len(rows)),
        ("Avg opportunity", round(sum(opps) / max(len(opps), 1), 1)),
        ("Classifications", ", ".join(f"{k.split('_')[0]} {v}" for k, v in sorted(by_cls.items()))),
        ("Sources", ", ".join(f"{k} {v}" for k, v in sorted(by_src.items()))),
    ]
    r = 6
    for label, val in stats:
        dash.cell(row=r, column=1, value=label)
        style_cell(dash.cell(row=r, column=1), bold=True, size=11)
        dash.cell(row=r, column=2, value=val)
        style_cell(dash.cell(row=r, column=2), size=11)
        dash.row_dimensions[r].height = 20
        r += 1
    dash.column_dimensions["A"].width = 22
    dash.column_dimensions["B"].width = 70

    # --- Sheet 2: leads ------------------------------------------------------
    leads = wb.create_sheet("Leads")
    if leads is None:
        raise RuntimeError('sheet missing')
    leads.sheet_view.showGridLines = False
    brand_block(leads, "Lead book", "One row per business. Every column traces to a stored CSV field.")
    headers = [
        "Business", "Industry", "Phone", "Website", "Classification",
        "Opportunity", "Need", "Value", "Site Health", "Rating", "Reviews",
        "Source", "Low Confidence",
    ]
    data = [
        [
            r["Business Name"], r["Industry"], r["Phone"], r["Website"], r["Classification"],
            num(r["Opportunity Score"]), num(r["Need Score"]), num(r["Value Score"]),
            r["Site Health"], r["Rating"], r["Review Count"], r["Sources"], r["Low Confidence"],
        ]
        for r in rows
    ]
    table(leads, headers, data, 6, [38, 16, 16, 34, 16, 11, 7, 7, 11, 8, 8, 10, 14])

    # --- Sheet 3: pipeline tracker (accounting/marketing) --------------------
    track = wb.create_sheet("Pipeline")
    if track is None:
        raise RuntimeError('sheet missing')
    track.sheet_view.showGridLines = False
    brand_block(
        track,
        "Outreach tracker",
        "Fill as you work. Blank cells by design -- log only what happened.",
    )
    headers = [
        "Business", "State (send/verify/expand/hold/drop)", "Channel (SMS/call/email)",
        "Date sent", "Replied?", "Reply date", "Meeting?", "Proposal sent?",
        "Closed?", "Revenue ($)", "Notes",
    ]
    data = [[r["Business Name"], "", "", "", "", "", "", "", "", "", ""] for r in rows]
    table(track, headers, data, 6, [34, 26, 20, 12, 10, 12, 10, 16, 8, 12, 40])

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
